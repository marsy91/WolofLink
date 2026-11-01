# Copyright (c) 2026 marsy91 (the WolofLink owner). All rights reserved. See LICENSE.
"""Builds the native-speaker review workbook.
Usage: python3 scripts/review/build-workbook.py out/WolofLink-Wolof-review.xlsx
Then copy public/audio/wo-ai/*.mp3 into an "audio" folder next to it and zip both."""
import json, re, sys, unicodedata, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

import os
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = sys.argv[1]
rows = json.loads(subprocess.check_output(["npx", "tsx", "scripts/review/export-lines.ts"], cwd=REPO))
holidays = json.loads(subprocess.check_output(
    ["npx", "tsx", "-e", "import {HOLIDAYS} from './src/content/calendar'; console.log(JSON.stringify(HOLIDAYS.map(h=>({id:h.id,title:h.title.en,titleFr:h.title.fr,about:h.about.en,aboutFr:h.about.fr,public:h.publicHoliday}))))"],
    cwd=REPO))

def norm(s):
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-zñŋ ]", "", s).split()

# 1. Audio the automatic checker flagged.
flagged = set(re.findall(r"`([A-Za-z0-9._:-]+)\.mp3`", open(f"{REPO}/docs/audio/ai-voice-report.md").read()))
# 2. Words and phrases the earlier review asked a native speaker to confirm.
review = open(f"{REPO}/docs/wolof-review.md").read()
asks = []
for block in re.findall(r"Native check, please:(.*?)(?:\n\n|\Z)", review, re.S):
    asks += re.findall(r"\*([^*]+)\*", block)
still = review.split("## Still needs a native speaker")[1].split("## Sources")[0]
asks += re.findall(r"\*([^*]+)\*", still)
ask_norm = [norm(a) for a in asks if len(norm(a)) > 0]
PRIORITY_CATS = {"Phrasebook: Health and emergencies", "Phrasebook: Feelings and opinions", "Phrasebook: Celebrations and faith"}

def reasons(r):
    why = []
    w = norm(r["wo"])
    if any(a == w or (len(a) >= 2 and " ".join(a) in " ".join(w)) for a in ask_norm):
        why.append("Spelling or usage not confirmed by sources")
    if r["id"] in flagged:
        why.append("Audio flagged by the automatic check")
    if r["where"] in PRIORITY_CATS:
        why.append("Health, feelings or faith phrase")
    return why

for r in rows:
    r["why"] = reasons(r)
    r["prio"] = 1 if r["why"] else 2
rows = sorted(rows, key=lambda r: r["prio"])

F = "Arial"
thin = Side(style="thin", color="D9CFC1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F2A4A")
head_font = Font(name=F, bold=True, color="FFFFFF", size=11)
input_fill = PatternFill("solid", fgColor="FFF6CC")
p1_fill = PatternFill("solid", fgColor="FCE4DC")
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ── Sheet 1: instructions ─────────────────────────────────────────────
ins = wb.active
ins.title = "How to review"
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 3
ins.column_dimensions["B"].width = 26
for c in "CDEFGH":
    ins.column_dimensions[c].width = 22
lines = [
    ("WolofLink: Wolof review", Font(name=F, bold=True, size=18, color="1F2A4A")),
    ("Thank you for helping! WolofLink is an app that teaches Senegalese Wolof (as spoken in Dakar) to English and French speakers.", None),
    ("", None),
    ("What to do", Font(name=F, bold=True, size=13, color="1F2A4A")),
    ("1. Open the “Lines to check” tab. Start with the rows marked Priority 1 (pink): they are the ones we are least sure about.", None),
    ("2. For each line, read the Wolof and the English/French meaning, and play the audio (click the file name; keep the “audio” folder next to this file).", None),
    ("3. Fill in ONLY the yellow columns: choose OK or Fix from each drop-down. When something needs fixing, write the correct version next to it.", None),
    ("4. Please use standard Senegalese (CLAD) spelling, as in Dakar: ë, é, à, ñ, ŋ, and double letters for long sounds (aa, ee, oo).", None),
    ("5. If a phrase is correct but nobody would say it that way in Dakar, choose Fix and write what people really say.", None),
    ("6. The “Holiday texts” tab has short descriptions of Senegalese holidays. Check they are accurate and respectful.", None),
    ("7. The “Progress” tab counts what you have finished. Send the file back when Priority 1 is done, then again when everything is done.", None),
    ("", None),
    ("Audio problems to listen for", Font(name=F, bold=True, size=13, color="1F2A4A")),
    ("Wrong word or missing word · wrong stress · short vowel that should be long (aa, ee, oo) or the reverse · ë, ñ or ŋ pronounced wrongly · sounds unnatural or robotic.", None),
    ("", None),
    ("Example of a filled-in row", Font(name=F, bold=True, size=13, color="1F2A4A")),
]
r = 1
for text, font in lines:
    c = ins.cell(row=r, column=2, value=text)
    c.font = font or Font(name=F, size=11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    if font is None and text:
        ins.row_dimensions[r].height = 30
    r += 1
ex_head = ["Wolof", "English", "Spelling OK?", "Correct Wolof", "Audio OK?", "Audio problem", "Comment"]
ex_vals = ["Jërëjëf", "Thank you", "OK", "", "Fix", "The last “ëf” is cut off", "Also common: Jërëjëf waay"]
for i, (h, v) in enumerate(zip(ex_head, ex_vals)):
    hc = ins.cell(row=r, column=2 + i, value=h)
    hc.font, hc.fill, hc.border = head_font, head_fill, border
    vc = ins.cell(row=r + 1, column=2 + i, value=v)
    vc.font, vc.border, vc.alignment = Font(name=F, size=11, italic=True), border, wrap
    if i >= 2:
        vc.fill = input_fill
r += 3
leg = ins.cell(row=r, column=2, value="Yellow cells = for you to fill in. Everything else is filled in already; please don't change it.")
leg.font = Font(name=F, size=11, bold=True)
leg.fill = input_fill
ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)

# ── Sheet 2: lines ────────────────────────────────────────────────────
ws = wb.create_sheet("Lines to check")
headers = [
    ("#", 5), ("Priority", 8), ("Why check first", 26), ("Wolof", 30), ("English meaning", 30), ("French meaning", 30),
    ("Note shown in the app", 30), ("Where in the app", 26), ("Audio (click to play)", 22),
    ("Spelling / wording OK?", 12), ("Correct Wolof", 28), ("Meaning OK?", 11), ("Better meaning", 26),
    ("Audio OK?", 10), ("Audio problem", 26), ("Comment", 30), ("ID (don't change)", 18),
]
for i, (h, w) in enumerate(headers, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.font, c.fill, c.border = head_font, head_fill, border
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 34
INPUT_COLS = [10, 11, 12, 13, 14, 15, 16]
for n, row in enumerate(rows, 1):
    rr = n + 1
    clip = row["audio"].split("/")[-1]
    vals = [n, row["prio"], "; ".join(row["why"]), row["wo"], row["en"], row["fr"], row["note"], row["where"],
            f"audio/{clip}", None, None, None, None, None, None, None, row["id"]]
    for ci, v in enumerate(vals, 1):
        c = ws.cell(row=rr, column=ci, value=v)
        c.font = Font(name=F, size=11, bold=(ci == 4))
        c.alignment = wrap
        c.border = border
        if ci in INPUT_COLS:
            c.fill = input_fill
        elif row["prio"] == 1 and ci <= 3:
            c.fill = p1_fill
    link = ws.cell(row=rr, column=9)
    link.hyperlink = f"audio/{clip}"
    link.font = Font(name=F, size=11, color="1F5FAD", underline="single")
last = len(rows) + 1
okfix = DataValidation(type="list", formula1='"OK,Fix"', allow_blank=True)
okfix.error = "Choose OK or Fix"
ws.add_data_validation(okfix)
for col in "JLN":
    okfix.add(f"{col}2:{col}{last}")
red = PatternFill("solid", fgColor="F8C9C4")
green = PatternFill("solid", fgColor="CDEBD3")
for col in "JLN":
    rng = f"{col}2:{col}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fix"'], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=green))
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"

# ── Sheet 3: holiday texts ────────────────────────────────────────────
hs = wb.create_sheet("Holiday texts")
hh = [("Holiday", 28), ("Holiday (French)", 28), ("Public holiday in the app", 12), ("Description (English)", 60),
      ("Description (French)", 60), ("Accurate?", 11), ("Correction or comment", 40)]
for i, (h, w) in enumerate(hh, 1):
    c = hs.cell(row=1, column=i, value=h)
    c.font, c.fill, c.border = head_font, head_fill, border
    c.alignment = Alignment(wrap_text=True, vertical="center")
    hs.column_dimensions[get_column_letter(i)].width = w
for n, h in enumerate(holidays, 2):
    vals = [h["title"], h["titleFr"], "Yes" if h["public"] else "No", h["about"], h["aboutFr"], None, None]
    for ci, v in enumerate(vals, 1):
        c = hs.cell(row=n, column=ci, value=v)
        c.font, c.alignment, c.border = Font(name=F, size=11, bold=(ci == 1)), wrap, border
        if ci >= 6:
            c.fill = input_fill
hlast = len(holidays) + 1
okfix2 = DataValidation(type="list", formula1='"OK,Fix"', allow_blank=True)
hs.add_data_validation(okfix2)
okfix2.add(f"F2:F{hlast}")
hs.conditional_formatting.add(f"F2:F{hlast}", CellIsRule(operator="equal", formula=['"Fix"'], fill=red))
hs.conditional_formatting.add(f"F2:F{hlast}", CellIsRule(operator="equal", formula=['"OK"'], fill=green))
hs.freeze_panes = "B2"

# ── Sheet 4: progress ─────────────────────────────────────────────────
ps = wb.create_sheet("Progress")
ps.sheet_view.showGridLines = False
ps.column_dimensions["A"].width = 3
ps.column_dimensions["B"].width = 44
for c in "CDE":
    ps.column_dimensions[c].width = 16
ps["B1"] = "Review progress"
ps["B1"].font = Font(name=F, bold=True, size=16, color="1F2A4A")
ps["B2"] = "Calculated automatically from the other tabs. A line counts as done when its three OK/Fix columns are all filled."
ps["B2"].font = Font(name=F, size=10, italic=True)
for i, h in enumerate(["", "Lines", "Done", "% done"]):
    if h:
        c = ps.cell(row=4, column=2 + i, value=h)
        c.font, c.fill, c.border = head_font, head_fill, border
L = f"'Lines to check'!"
rngB, rngJ, rngL, rngN = (f"{L}${c}$2:${c}${last}" for c in "BJLN")
table = [
    ("Priority 1 (check first)", f"=COUNTIF({rngB},1)", f'=COUNTIFS({rngB},1,{rngJ},"<>",{rngL},"<>",{rngN},"<>")'),
    ("Priority 2", f"=COUNTIF({rngB},2)", f'=COUNTIFS({rngB},2,{rngJ},"<>",{rngL},"<>",{rngN},"<>")'),
    ("All lines", "=C5+C6", "=D5+D6"),
]
for i, (label, total, done) in enumerate(table, 5):
    ps.cell(row=i, column=2, value=label).font = Font(name=F, size=11, bold=(i == 7))
    ps.cell(row=i, column=3, value=total).font = Font(name=F, size=11)
    ps.cell(row=i, column=4, value=done).font = Font(name=F, size=11)
    pc = ps.cell(row=i, column=5, value=f"=IF(C{i}=0,0,D{i}/C{i})")
    pc.number_format = "0%"
    pc.font = Font(name=F, size=11, bold=True)
    for c in range(2, 6):
        ps.cell(row=i, column=c).border = border
ps["B9"] = "Marked “Fix”"
ps["B9"].font = Font(name=F, bold=True, size=12, color="1F2A4A")
fixes = [
    ("Spelling / wording to fix", f'=COUNTIF({rngJ},"Fix")'),
    ("Meanings to fix", f'=COUNTIF({rngL},"Fix")'),
    ("Audio to fix", f'=COUNTIF({rngN},"Fix")'),
    ("Holiday texts to fix", f"=COUNTIF('Holiday texts'!$F$2:$F${hlast},\"Fix\")"),
]
for i, (label, f) in enumerate(fixes, 10):
    ps.cell(row=i, column=2, value=label).font = Font(name=F, size=11)
    c = ps.cell(row=i, column=3, value=f)
    c.font = Font(name=F, size=11)
    for col in (2, 3):
        ps.cell(row=i, column=col).border = border
ps["B15"] = f"Priority 1 = {sum(1 for x in rows if x['prio']==1)} lines whose spelling was not confirmed by sources, whose audio was flagged by the automatic check, or that are health, feelings or faith phrases."
ps["B15"].font = Font(name=F, size=10, italic=True)
ps["B15"].alignment = Alignment(wrap_text=True)
ps.merge_cells("B15:E16")
ps.row_dimensions[15].height = 30

wb.move_sheet("Progress", offset=-2)
wb.active = 0
wb.save(OUT)
print("rows", len(rows), "priority1", sum(1 for x in rows if x["prio"] == 1), "asks", len(asks), "flagged", len(flagged))
